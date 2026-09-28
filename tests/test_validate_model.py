"""Tests for validate_model.py, by breaking a copy of the model on purpose.

A validator that only ever runs against a correct workbook proves nothing. It
would report "all checks pass" just as confidently if its comparisons were
inverted, its tolerance were infinite, or it were reading the wrong rows. The
only way to know it works is to hand it a workbook that is wrong and confirm it
says so.

So each test copies the real Apple model, breaks one specific thing, and asserts
the validator catches that thing. Four failure modes are covered:

* **Arithmetic that does not add up.** The case the script exists for.
* **A duplicate row label.** The validator keys rows by their label. Two rows
  sharing a label would silently bind every check to whichever came first and
  report a pass for a row nobody chose, which is worse than not checking.
* **Statements laid out over different years.** Comparisons line up by position,
  so mismatched headers would compare FY21 against FY20 and call it agreement.
* **A Validation figure that no longer matches its source statement.** The
  validator reads the Validation tab, which is a transcription. Until the
  transcription checks existed, wiping every numeric cell on the IS and BS still
  produced "19 of 19 matched" and exit 0.

The duplicate-label and year-header cases are structural. They hold in the Apple
model today, but the script is documented as working on any workbook with a
Validation tab, so they are properties to enforce rather than facts to rely on.
"""
import os
import shutil
import stat
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

import xlsx_surgery

REPO = Path(__file__).resolve().parent.parent
MODEL = REPO / "Apple" / "Apple_Financial_Model.xlsx"
SCRIPT = REPO / "validate_model.py"

# Rows in the Validation tab, used to inject faults.
ROW_TOTAL_REVENUE = 18
ROW_BS_HEADER = 35


def _run(workbook: Path):
    """Run the validator on a workbook, returning (exit code, combined output)."""
    result = subprocess.run(
        [sys.executable, str(SCRIPT), str(workbook)],
        capture_output=True,
        text=True,
        cwd=REPO,
    )
    return result.returncode, result.stdout + result.stderr


@pytest.fixture
def workbook_copy(tmp_path):
    """A writable copy of the real model, for one test to damage."""

    def _copy(name="model.xlsx"):
        dst = tmp_path / name
        # copyfile rather than copy: the source may be read-only and copy would
        # carry that across, leaving openpyxl unable to save.
        shutil.copyfile(MODEL, dst)
        os.chmod(dst, stat.S_IWRITE | stat.S_IREAD)
        return dst

    return _copy


def _edit(path: Path, edits):
    """Structural edits, through openpyxl.

    Saving with openpyxl discards every cached formula result in the file, so
    this can only be used for tests whose expected failure is raised before the
    transcription checks read IS, BS and CFS. Use `_set_number` for anything
    that has to leave the rest of the workbook readable.
    """
    wb = openpyxl.load_workbook(path)
    ws = wb["Validation"]
    for (row, col), value in edits.items():
        ws.cell(row, col).value = value
    wb.save(path)
    wb.close()


def _set_number(path: Path, row: int, col: int, value):
    """Change one number on the Validation tab and nothing else.

    The validator now reads IS, BS and CFS as well, and those are formula cells
    whose values only exist in the file as cached results. openpyxl would drop
    all of them on save, so every transcription check would fail with "missing
    value" and the test would be measuring the save, not the fault.
    """
    xlsx_surgery.set_cached_value(path, "Validation", xlsx_surgery.ref(row, col), value)


# --------------------------------------------------------------------------
# The happy path, so a failure elsewhere is not just "the script is broken"
# --------------------------------------------------------------------------

def test_the_real_model_passes():
    """The shipped model's identities all hold. If this fails, either the model
    changed or the validator did, and the other tests cannot be interpreted."""
    code, out = _run(MODEL)
    assert code == 0, out
    assert "accounting identities re-derived independently and matched" in out


def test_advisory_does_not_fail_the_run():
    """The cash-flow-to-balance-sheet comparison depends on whether a filing
    reconciles to cash including restricted cash. That is a presentation
    question, so it is reported and must not fail the build."""
    code, out = _run(MODEL)
    assert code == 0
    assert "NOTE" in out
    assert "advisory" in out


# --------------------------------------------------------------------------
# Arithmetic
# --------------------------------------------------------------------------

def test_broken_total_is_caught(workbook_copy):
    """Change total revenue so it no longer equals its two segments. The
    validator must fail, and must fail loudly enough to exit non-zero."""
    wb = workbook_copy()
    _set_number(wb, ROW_TOTAL_REVENUE, 2, 1)
    code, out = _run(wb)
    assert code == 1, out
    assert "FAIL" in out
    assert "revenue = products + services" in out


def test_failure_names_the_year_and_the_gap(workbook_copy):
    """A failure that does not say which year and by how much sends the reader
    back to the spreadsheet to find it themselves."""
    wb = workbook_copy()
    _set_number(wb, ROW_TOTAL_REVENUE, 2, 1)
    _, out = _run(wb)
    assert "FY21A" in out
    assert "differs by" in out


def test_a_small_rounding_difference_still_passes(workbook_copy):
    """Figures are whole millions on values in the hundreds of billions, so an
    exact-equality test would fail on rounding alone. One million of slack is
    deliberate, and it has to actually be there."""
    wb = workbook_copy()
    original = openpyxl.load_workbook(MODEL, data_only=True)["Validation"].cell(ROW_TOTAL_REVENUE, 2).value
    _set_number(wb, ROW_TOTAL_REVENUE, 2, original + 0.4)
    code, out = _run(wb)
    assert code == 0, out


# --------------------------------------------------------------------------
# Structure
# --------------------------------------------------------------------------

def test_duplicate_row_label_is_rejected(workbook_copy):
    """Two rows sharing a label would make every check on it validate whichever
    came first and report a pass. The validator must refuse rather than pick."""
    wb = workbook_copy()
    book = openpyxl.load_workbook(wb)
    ws = book["Validation"]
    row = ws.max_row + 2
    ws.cell(row, 1).value = "Total revenue (calc)"
    for col in range(2, 7):
        ws.cell(row, col).value = 999
    book.save(wb)
    book.close()

    code, out = _run(wb)
    assert code == 1, out
    assert "Duplicate numeric row labels" in out


def test_a_statement_that_no_longer_matches_the_transcription_is_caught(workbook_copy):
    """The gap this validator had until the transcription checks existed.

    Every identity it checks is computed from the Validation tab, which is a
    copy of the statements. Change the statement and leave the copy alone and
    all of them still pass, because none of them ever opened the statement.
    Here FY21A revenue on the IS is changed and nothing on the Validation tab
    is, so every identity still holds and the run must still fail.
    """
    wb = workbook_copy()
    xlsx_surgery.set_cached_value(wb, "IS", "B8", 1)
    code, out = _run(wb)
    assert code == 1, out
    assert "SRC Products revenue" in out
    assert "19 of 19 accounting identities" in out, (
        "the identities are expected to still pass; that is the point of the test"
    )


def test_mismatched_year_headers_are_rejected(workbook_copy):
    """Comparisons line up by position. If the balance sheet ran over different
    years, every check would still "pass" while comparing the wrong columns."""
    wb = workbook_copy()
    _edit(wb, {(ROW_BS_HEADER, 2): "FY20A"})
    code, out = _run(wb)
    assert code == 1, out
    assert "not laid out over the same years" in out


def test_missing_validation_sheet_is_reported_clearly(workbook_copy, tmp_path):
    """A workbook without a Validation tab should say so and list what it does
    have, rather than raising a KeyError from deep inside openpyxl."""
    wb = workbook_copy()
    book = openpyxl.load_workbook(wb)
    del book["Validation"]
    book.save(wb)
    book.close()

    code, out = _run(wb)
    assert code != 0
    assert "no Validation sheet" in out


def test_missing_file_is_reported_clearly(tmp_path):
    code, out = _run(tmp_path / "does_not_exist.xlsx")
    assert code != 0
    assert "workbook not found" in out
